"""
Feature to Excel - single-file offline desktop tool

Requirements:
    pip install openpyxl

Run:
    python feature_to_excel.py

What it does:
- Select one .feature file
- Parses Feature, tags, Scenario, Scenario Outline, steps and Examples
- Expands every Examples row into a separate Excel test case
- Generates columns:
  ID, Labels, Lan, Test Summary, Description, Action, Expected Result,
  Data, Auto Assessment, POD, MALCODE, Test Repository Path
"""

from __future__ import annotations

import re
import tkinter as tk
from dataclasses import dataclass, field
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


# -----------------------------
# Models
# -----------------------------

@dataclass
class Step:
    keyword: str
    text: str


@dataclass
class ExampleRow:
    values: Dict[str, str]
    tags: List[str] = field(default_factory=list)


@dataclass
class Scenario:
    name: str
    tags: List[str] = field(default_factory=list)
    steps: List[Step] = field(default_factory=list)
    examples: List[ExampleRow] = field(default_factory=list)


@dataclass
class Feature:
    name: str
    tags: List[str] = field(default_factory=list)
    scenarios: List[Scenario] = field(default_factory=list)


@dataclass
class ExcelTestCase:
    test_id: str
    labels: str
    language: str
    test_summary: str
    description: str
    action: str
    expected_result: str
    data: str = ""
    auto_assessment: str = "A: In Sprint Automation"
    pod: str = "Batman"
    malcode: str = "HOJ"
    test_repository_path: str = ""


class FeatureParseError(ValueError):
    pass


# -----------------------------
# Gherkin parser
# -----------------------------

class FeatureParser:
    FEATURE_RE = re.compile(r"^\s*Feature\s*:\s*(.+)$", re.IGNORECASE)
    SCENARIO_RE = re.compile(
        r"^\s*(Scenario(?:\s+Outline)?|Scenario Template)\s*:\s*(.+)$",
        re.IGNORECASE,
    )
    STEP_RE = re.compile(r"^\s*(Given|When|Then|And|But|\*)\s+(.+)$", re.IGNORECASE)
    EXAMPLES_RE = re.compile(r"^\s*Examples\s*:\s*$", re.IGNORECASE)

    def parse_file(self, path: str | Path) -> Feature:
        feature_path = Path(path)
        if not feature_path.exists():
            raise FileNotFoundError(f"Feature file not found: {feature_path}")
        return self.parse_text(feature_path.read_text(encoding="utf-8-sig"))

    def parse_text(self, text: str) -> Feature:
        lines = text.splitlines()
        feature: Feature | None = None
        current_scenario: Scenario | None = None
        pending_tags: List[str] = []
        i = 0

        while i < len(lines):
            raw = lines[i]
            line = raw.strip()

            if not line or line.startswith("#"):
                i += 1
                continue

            if line.startswith("@"):
                pending_tags.extend(line.split())
                i += 1
                continue

            feature_match = self.FEATURE_RE.match(raw)
            if feature_match:
                feature = Feature(
                    name=feature_match.group(1).strip(),
                    tags=pending_tags.copy(),
                )
                pending_tags.clear()
                i += 1
                continue

            scenario_match = self.SCENARIO_RE.match(raw)
            if scenario_match:
                if feature is None:
                    raise FeatureParseError("Scenario found before Feature declaration.")

                current_scenario = Scenario(
                    name=scenario_match.group(2).strip(),
                    tags=pending_tags.copy(),
                )
                pending_tags.clear()
                feature.scenarios.append(current_scenario)
                i += 1
                continue

            step_match = self.STEP_RE.match(raw)
            if step_match and current_scenario:
                current_scenario.steps.append(
                    Step(
                        keyword=step_match.group(1).title(),
                        text=step_match.group(2).strip(),
                    )
                )
                i += 1
                continue

            if self.EXAMPLES_RE.match(raw) and current_scenario:
                example_tags = pending_tags.copy()
                pending_tags.clear()
                i = self._parse_examples(lines, i + 1, current_scenario, example_tags)
                continue

            i += 1

        if feature is None:
            raise FeatureParseError("No Feature declaration was found.")
        if not feature.scenarios:
            raise FeatureParseError("No Scenario or Scenario Outline was found.")

        return feature

    def _parse_examples(
        self,
        lines: List[str],
        start_index: int,
        scenario: Scenario,
        example_tags: List[str],
    ) -> int:
        i = start_index

        while i < len(lines):
            stripped = lines[i].strip()

            if not stripped or stripped.startswith("#"):
                i += 1
                continue

            if stripped.startswith("@") or self.SCENARIO_RE.match(lines[i]):
                return i

            if stripped.startswith("|"):
                break

            i += 1

        if i >= len(lines) or not lines[i].strip().startswith("|"):
            return i

        headers = self._split_table_row(lines[i])
        i += 1

        while i < len(lines):
            stripped = lines[i].strip()

            if not stripped:
                i += 1
                continue

            if not stripped.startswith("|"):
                break

            values = self._split_table_row(lines[i])

            if len(values) != len(headers):
                raise FeatureParseError(
                    f"Examples row has {len(values)} values but {len(headers)} headers:\n"
                    f"{lines[i]}"
                )

            scenario.examples.append(
                ExampleRow(
                    values=dict(zip(headers, values)),
                    tags=example_tags.copy(),
                )
            )
            i += 1

        return i

    @staticmethod
    def _split_table_row(row: str) -> List[str]:
        return [cell.strip().replace("\\n", "\n") for cell in row.strip().strip("|").split("|")]


# -----------------------------
# Feature -> test-case conversion
# -----------------------------

class TestCaseConverter:
    PLACEHOLDER_RE = re.compile(r"<([^>]+)>")

    def convert(self, feature: Feature) -> List[ExcelTestCase]:
        test_cases: List[ExcelTestCase] = []
        counter = 1

        for scenario in feature.scenarios:
            example_rows = scenario.examples or [ExampleRow(values={})]

            for example_row in example_rows:
                example = example_row.values
                resolved_name = self._replace_placeholders(scenario.name, example)
                resolved_steps = [
                    (step.keyword, self._replace_placeholders(step.text, example))
                    for step in scenario.steps
                ]

                action_lines: List[str] = []
                expected_lines: List[str] = []
                in_expected_result = False

                for keyword, text in resolved_steps:
                    normalized_keyword = keyword.lower()

                    if normalized_keyword in {"given", "when"}:
                        in_expected_result = False
                    elif normalized_keyword == "then":
                        in_expected_result = True

                    line = f"{keyword.upper()} {text}"

                    if in_expected_result:
                        expected_lines.append(line)
                    else:
                        action_lines.append(line)

                all_tags = self._unique(feature.tags + scenario.tags + example_row.tags)
                labels = self._value_for(example, {"label", "labels"})
                if not labels:
                    labels = ", ".join(self._extract_labels(all_tags))
                language = self._detect_language(all_tags, example)
                test_summary = self._build_test_summary(resolved_name, example, language)

                test_cases.append(
                    ExcelTestCase(
                        test_id=f"TC-{counter:03d}",
                        labels=labels,
                        language=language,
                        test_summary=test_summary,
                        description=self._build_description(test_summary, example, language),
                        action="\n".join(action_lines),
                        expected_result="\n".join(expected_lines),
                        data=self._value_for(example, {"data"}),
                        auto_assessment=(
                            self._value_for(example, {"autoassessment", "autoassessmentstatus"})
                            or "A: In Sprint Automation"
                        ),
                        pod=self._value_for(example, {"pod"}) or "Batman",
                        malcode=self._value_for(example, {"malcode", "malcodevalue"}) or "HOJ",
                        test_repository_path=self._value_for(
                            example,
                            {"testrepositorypath", "repositorypath", "testrepopath"},
                        ),
                    )
                )
                counter += 1

        return test_cases

    def _replace_placeholders(self, text: str, values: Dict[str, str]) -> str:
        def replacement(match: re.Match[str]) -> str:
            key = match.group(1)
            return values.get(key, match.group(0))

        return self.PLACEHOLDER_RE.sub(replacement, text)

    def _build_test_summary(
        self,
        fallback_summary: str,
        example: Dict[str, str],
        language: str,
    ) -> str:
        explicit_summary = self._value_for(example, {"testsummary"})
        if explicit_summary:
            return explicit_summary

        application = self._value_for(example, {"application", "app"})
        user_role = self._value_for(example, {"userrole", "role"})
        if not user_role:
            user_role = self._user_role_from_sso(
                self._value_for(example, {"ssouser", "user"})
            )

        parts = [
            application,
            language,
            user_role,
            self._value_for(example, {"page"}),
            self._value_for(example, {"mainfields", "productfields"}),
            self._value_for(example, {"notificationtype"}),
            self._value_for(example, {"notificationname"}),
            self._value_for(example, {"state", "status"}),
        ]
        summary_parts = [part for part in parts if part]

        if len(summary_parts) >= 4:
            return "_".join(summary_parts)

        return fallback_summary

    def _build_description(
        self,
        summary: str,
        example: Dict[str, str],
        language: str,
    ) -> str:
        explicit_description = self._value_for(example, {"description"})
        if explicit_description:
            return explicit_description

        notification_text = self._value_for(
            example,
            {"notificationtext", "notifications"},
        )
        notification_category = self._value_for(
            example,
            {"notificationcategory", "category"},
        )
        user_display = self._value_for(example, {"userdisplay"})
        if not user_display:
            user_display = f"Retail Advisor ({language})"
        elif f"({language})" not in user_display:
            user_display = f"{user_display} ({language})"

        if notification_text and notification_category:
            return (
                f"Verify that a {user_display} can view the notification\n"
                f"{notification_text}\n"
                f"under correct notification category ({notification_category}) "
                "on RB Home Page"
            )

        if not summary:
            return "Verify that the scenario works as expected."
        return f"Verify that {summary[0].lower() + summary[1:]}."

    @staticmethod
    def _value_for(values: Dict[str, str], keys: set[str]) -> str:
        for key, value in values.items():
            normalized_key = re.sub(r"[^a-z0-9]", "", key.lower())
            if normalized_key in keys:
                return value.strip()
        return ""

    @staticmethod
    def _user_role_from_sso(sso_user: str) -> str:
        match = re.search(r"RETAIL_([A-Z]+)UserId", sso_user, re.IGNORECASE)
        if match:
            return match.group(1).upper()
        return ""

    @staticmethod
    def _unique(items: Iterable[str]) -> List[str]:
        result: List[str] = []
        seen = set()

        for item in items:
            if item not in seen:
                result.append(item)
                seen.add(item)

        return result

    @staticmethod
    def _extract_labels(tags: List[str]) -> List[str]:
        labels: List[str] = []

        for tag in tags:
            cleaned = tag.lstrip("@")
            lowered = cleaned.lower()

            if lowered.endswith("_en") or lowered.endswith("_fr"):
                cleaned = cleaned[:-3]

            if cleaned and cleaned.lower() not in {"en", "fr", "english", "french"}:
                labels.append(cleaned)

        return labels

    @staticmethod
    def _detect_language(tags: List[str], example: Dict[str, str]) -> str:
        for key, value in example.items():
            if key.lower() in {"lan", "lang", "language"}:
                normalized = value.strip().upper()
                if normalized.startswith("FR"):
                    return "FR"
                if normalized.startswith("EN"):
                    return "EN"

        for tag in tags:
            lowered = tag.lower()
            if lowered.endswith("_fr") or lowered in {"@fr", "@french"}:
                return "FR"
            if lowered.endswith("_en") or lowered in {"@en", "@english"}:
                return "EN"

        return "EN"


# -----------------------------
# Excel generation
# -----------------------------

class ExcelGenerator:
    HEADERS = [
        "ID",
        "Labels",
        "Lan",
        "Test Summary",
        "Description",
        "Action",
        "Expected Result",
        "Data",
        "Auto Assessment",
        "POD",
        "MALCODE",
        "Test Repository Path",
    ]

    HEADER_FILL = "FFF200"
    EN_FILL = "FFF4CC"
    FR_FILL = "EADCF8"
    BORDER_COLOR = "B7B7B7"

    def generate(
        self,
        test_cases: List[ExcelTestCase],
        output_path: str | Path,
    ) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Test Cases"

        sheet.append(self.HEADERS)

        for test_case in test_cases:
            sheet.append(
                [
                    test_case.test_id,
                    test_case.labels,
                    test_case.language,
                    test_case.test_summary,
                    test_case.description,
                    test_case.action,
                    test_case.expected_result,
                    test_case.data,
                    test_case.auto_assessment,
                    test_case.pod,
                    test_case.malcode,
                    test_case.test_repository_path,
                ]
            )

        self._format_sheet(sheet, len(test_cases) + 1)
        workbook.save(output_path)
        return output_path

    def _format_sheet(self, sheet, last_row: int) -> None:
        thin_side = Side(style="thin", color=self.BORDER_COLOR)
        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=self.HEADER_FILL)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = border

        for row_number in range(2, last_row + 1):
            language = str(sheet.cell(row=row_number, column=3).value or "").upper()
            row_fill = self.FR_FILL if language == "FR" else self.EN_FILL

            for column_number in range(1, len(self.HEADERS) + 1):
                cell = sheet.cell(row=row_number, column=column_number)
                if column_number <= 7:
                    cell.fill = PatternFill("solid", fgColor=row_fill)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

        column_widths = {
            "A": 12,
            "B": 30,
            "C": 8,
            "D": 45,
            "E": 55,
            "F": 75,
            "G": 75,
            "H": 16,
            "I": 18,
            "J": 14,
            "K": 14,
            "L": 24,
        }

        for column, width in column_widths.items():
            sheet.column_dimensions[column].width = width

        sheet.row_dimensions[1].height = 28

        for row_number in range(2, last_row + 1):
            sheet.row_dimensions[row_number].height = 115

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:L{last_row}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1

        table = Table(
            displayName="GeneratedTestCases",
            ref=f"A1:L{last_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        sheet.add_table(table)


# -----------------------------
# Desktop UI
# -----------------------------

class FeatureToExcelApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()

        self.title("Feature to Excel")
        self.geometry("780x350")
        self.minsize(700, 320)

        self.feature_path = tk.StringVar()
        self.output_path = tk.StringVar()

        self._build_ui()

    def _build_ui(self) -> None:
        container = ttk.Frame(self, padding=20)
        container.pack(fill="both", expand=True)

        ttk.Label(
            container,
            text="Feature File to Excel Test Cases",
            font=("Segoe UI", 16, "bold"),
        ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 18))

        ttk.Label(container, text="Feature file:").grid(
            row=1,
            column=0,
            sticky="w",
            pady=8,
        )

        ttk.Entry(
            container,
            textvariable=self.feature_path,
            width=70,
        ).grid(row=1, column=1, sticky="ew", padx=10)

        ttk.Button(
            container,
            text="Browse",
            command=self.select_feature,
        ).grid(row=1, column=2)

        ttk.Label(container, text="Output Excel:").grid(
            row=2,
            column=0,
            sticky="w",
            pady=8,
        )

        ttk.Entry(
            container,
            textvariable=self.output_path,
            width=70,
        ).grid(row=2, column=1, sticky="ew", padx=10)

        ttk.Button(
            container,
            text="Browse",
            command=self.select_output,
        ).grid(row=2, column=2)

        self.status_label = ttk.Label(
            container,
            text="Select a .feature file and an output location.",
            foreground="#555555",
        )
        self.status_label.grid(
            row=3,
            column=0,
            columnspan=3,
            sticky="w",
            pady=(18, 10),
        )

        ttk.Button(
            container,
            text="Generate Excel",
            command=self.generate_excel,
        ).grid(row=4, column=0, columnspan=3, pady=12)

        container.columnconfigure(1, weight=1)

    def select_feature(self) -> None:
        selected = filedialog.askopenfilename(
            title="Select feature file",
            filetypes=[
                ("Gherkin feature files", "*.feature"),
                ("All files", "*.*"),
            ],
        )

        if selected:
            self.feature_path.set(selected)

            if not self.output_path.get():
                feature_file = Path(selected)
                default_output = feature_file.with_name(
                    feature_file.stem + "_TestCases.xlsx"
                )
                self.output_path.set(str(default_output))

    def select_output(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="Save generated Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel workbook", "*.xlsx")],
        )

        if selected:
            self.output_path.set(selected)

    def generate_excel(self) -> None:
        feature_file = self.feature_path.get().strip()
        output_file = self.output_path.get().strip()

        if not feature_file:
            messagebox.showerror(
                "Missing feature file",
                "Please select a .feature file.",
            )
            return

        if not output_file:
            messagebox.showerror(
                "Missing output file",
                "Please choose the output Excel location.",
            )
            return

        try:
            self.status_label.config(text="Reading and converting feature file...")
            self.update_idletasks()

            feature = FeatureParser().parse_file(feature_file)
            test_cases = TestCaseConverter().convert(feature)
            generated_file = ExcelGenerator().generate(test_cases, output_file)

            self.status_label.config(
                text=f"Completed: {len(test_cases)} test case(s) generated."
            )

            messagebox.showinfo(
                "Excel generated",
                f"Generated {len(test_cases)} test case(s).\n\n{generated_file}",
            )

        except Exception as error:
            self.status_label.config(text="Generation failed.")
            messagebox.showerror("Generation failed", str(error))


if __name__ == "__main__":
    FeatureToExcelApp().mainloop()
